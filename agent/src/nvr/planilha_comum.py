"""Montagem das planilhas de configuracao dos NVR (Motorola, Hikvision, Intelbras): estilo, abas e a aba
'Câmeras (completo)', que junta numa linha so todos os dados de cada camera, cada um na sua coluna."""
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COR_BASE = "1F4E78"
COR_STREAM = "7030A0"
COR_AGENDA = "2E7D32"
COR_MOVIMENTO = "C55A11"
COR_EVENTOS = "A61C3C"
COR_IMAGEM = "00796B"

# Caracteres que o Excel nao aceita em celulas
_ILEGAIS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

# Campos que nunca vao para uma planilha (usuario e senha de cameras, DDNS, e-mail etc.)
_SENSIVEIS = {"password", "passwd", "pwd", "username", "secretkey"}


def eh_sensivel(nome) -> bool:
    """True para tag/atributo/chave que guarda usuario ou senha (compara o ultimo trecho do caminho)."""
    n = re.sub(r"\[.*?\]", "", str(nome)).split(".")[-1].lower()
    return n in _SENSIVEIS or n.endswith(("password", "passwd"))


# Atributos SMART (padrao ATA) mais comuns, para a aba SMART
SMART_NOMES = {
    "1": "Raw Read Error Rate", "3": "Spin-Up Time", "4": "Start/Stop Count", "5": "Reallocated Sectors",
    "7": "Seek Error Rate", "9": "Power-On Hours", "10": "Spin Retry Count", "11": "Calibration Retry Count",
    "12": "Power Cycle Count", "183": "SATA Downshift Error Count", "184": "End-to-End Error",
    "187": "Reported Uncorrectable", "188": "Command Timeout", "189": "High Fly Writes",
    "190": "Airflow Temperature", "191": "G-Sense Error Rate", "192": "Power-Off Retract Count",
    "193": "Load Cycle Count", "194": "Temperature", "195": "Hardware ECC Recovered",
    "196": "Reallocation Events", "197": "Current Pending Sectors", "198": "Offline Uncorrectable",
    "199": "UDMA CRC Errors", "200": "Multi-Zone Error Rate", "240": "Head Flying Hours",
    "241": "Total LBAs Written", "242": "Total LBAs Read",
}


def novo_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def estilizar(ws, largura_max=70):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=COR_BASE)
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    for col in range(1, ws.max_column + 1):
        mais_largo = max(len(str(c.value)) if c.value is not None else 0 for c in ws[get_column_letter(col)])
        ws.column_dimensions[get_column_letter(col)].width = min(max(mais_largo + 2, 10), largura_max)


def aba(wb, titulo, cabecalho, linhas):
    ws = wb.create_sheet(titulo)
    ws.append(cabecalho)
    for linha in linhas:
        ws.append([_ILEGAIS.sub("", v) if isinstance(v, str) else v for v in linha])
    if not linhas:
        ws.append(["(sem dados neste NVR)"] + [""] * (len(cabecalho) - 1))
    estilizar(ws)
    return ws


def aba_resumo(wb, campos):
    """Aba 'Resumo' (Campo/Valor), sempre a primeira."""
    ws = wb.create_sheet("Resumo")
    ws.append(["Campo", "Valor"])
    for campo, valor in campos:
        ws.append([campo, _ILEGAIS.sub("", valor) if isinstance(valor, str) else valor])
    estilizar(ws)
    return ws


def consolidar_cameras(wb, grupos, base="Câmeras", id_final="ID do canal"):
    """Cria a aba 'Câmeras (completo)' logo depois do Resumo e a deixa como aba ativa.

    `base`: aba com uma linha por camera (primeira coluna = canal). `grupos`: [(prefixo, nome da aba, cor)];
    cada aba de grupo tem canal e nome nas duas primeiras colunas e os demais dados em colunas separadas.
    A coluna `id_final` (se existir na base) vai para o fim."""
    def linhas(nome):
        todas = list(wb[nome].iter_rows(values_only=True))
        return todas[0], {str(r[0]): r for r in todas[1:] if r[0] is not None and not str(r[0]).startswith("(sem")}

    cab_base, dados_base = linhas(base)
    idx_id = cab_base.index(id_final) if id_final in cab_base else None
    colunas = [i for i in range(len(cab_base)) if i != idx_id]
    cabecalho = [cab_base[i] for i in colunas]
    fontes, cores = [], {}
    for prefixo, nome, cor in grupos:
        cab, dados = linhas(nome)
        cabecalho += [prefixo + c for c in cab[2:]]
        fontes.append((len(cab) - 2, dados))
        cores[prefixo] = cor
    if idx_id is not None:
        cabecalho.append(id_final)

    saida = []
    for canal, linha_base in dados_base.items():
        linha = [linha_base[i] for i in colunas]
        for tamanho, dados in fontes:
            origem = dados.get(canal)
            linha += list(origem[2:]) if origem else [""] * tamanho
        if idx_id is not None:
            linha.append(linha_base[idx_id])
        saida.append(linha)

    ws = aba(wb, "Câmeras (completo)", cabecalho, saida)
    ws.freeze_panes = "C2"
    for c in ws[1]:
        for prefixo, cor in cores.items():
            if str(c.value).startswith(prefixo):
                c.fill = PatternFill("solid", fgColor=cor)
    wb.move_sheet(ws, offset=-(len(wb.sheetnames) - 2))  # logo depois do Resumo
    for outra in wb.worksheets:  # abre nesta aba, e so nela (senao o Excel agrupa abas)
        outra.sheet_view.tabSelected = outra is ws
    wb.active = wb.worksheets.index(ws)
    return ws
